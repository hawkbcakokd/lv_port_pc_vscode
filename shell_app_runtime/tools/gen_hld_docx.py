# -*- coding: utf-8 -*-
# Chinese text loaded from hld_zh.json (UTF-8).
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    if bold:
        run.bold = True
    p.paragraph_format.space_after = Pt(6)


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (a, b) in enumerate(rows):
        t.rows[i].cells[0].text = a
        t.rows[i].cells[1].text = b
    doc.add_paragraph()


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    root = os.path.abspath(os.path.join(here, "..", ".."))
    out = os.path.join(root, "\u8f6f\u4ef6\u6982\u8981\u8bbe\u8ba1.docx")
    path = os.path.join(here, "hld_zh.json")
    raw = open(path, "rb").read()
    data = None
    for enc in ("utf-8", "utf-8-sig", "gbk"):
        try:
            data = json.loads(raw.decode(enc))
            break
        except (UnicodeDecodeError, json.JSONDecodeError):
            continue
    if data is None:
        raise RuntimeError("hld_zh.json decode or parse failed")

    doc = Document()
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(data["title"])
    r0.bold = True
    r0.font.size = Pt(18)
    doc.add_paragraph()
    for line in data["meta"]:
        add_body(doc, line)

    for sec in data["sections"]:
        doc.add_heading(sec["h1"], level=1)
        for sub in sec.get("subs") or []:
            doc.add_heading(sub["h2"], level=2)
            for para in sub.get("paras") or []:
                add_body(doc, para)
        bpi = sec.get("bold_para_index")
        for ip, para in enumerate(sec.get("paras") or []):
            add_body(doc, para, bold=(bpi is not None and ip == bpi))
        for b in sec.get("bullets") or []:
            add_body(doc, b)
        if sec.get("table"):
            add_table(doc, sec["table"])
        c = sec.get("conclusion")
        if c:
            add_body(doc, "\u7ed3\u8bba\uff1a" + c, bold=True)
        for para in sec.get("tail_paras") or []:
            add_body(doc, para)

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
